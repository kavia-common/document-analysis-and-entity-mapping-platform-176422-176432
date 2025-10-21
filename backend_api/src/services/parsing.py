from __future__ import annotations

import io
from pathlib import Path
from typing import Tuple

from PIL import Image
from pdfminer.high_level import extract_text as pdf_extract_text
from docx import Document as DocxDocument  # type: ignore
from pptx import Presentation  # type: ignore
from openpyxl import load_workbook  # type: ignore

from src.core.logger import get_logger
from .ocr import ocr_image

logger = get_logger(__name__)


# PUBLIC_INTERFACE
def parse_file_to_text(path: str) -> Tuple[str, dict]:
    """Parse a supported file into plain text and provide simple metadata."""
    p = Path(path)
    suffix = p.suffix.lower()
    meta: dict = {"name": p.name, "suffix": suffix}

    try:
        if suffix == ".pdf":
            text = pdf_extract_text(path) or ""
            meta["source"] = "pdf"
            return text, meta
        if suffix in (".docx",):
            doc = DocxDocument(path)
            text = "\n".join([p.text or "" for p in doc.paragraphs])
            meta["source"] = "docx"
            return text, meta
        if suffix in (".pptx",):
            prs = Presentation(path)
            slides_text = []
            for slide in prs.slides:
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text)
                slides_text.append("\n".join(slide_text))
            text = "\n\n".join(slides_text)
            meta["source"] = "pptx"
            return text, meta
        if suffix in (".xlsx", ".xlsm"):
            wb = load_workbook(path, read_only=True, data_only=True)
            sheet_texts = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v) if v is not None else "" for v in row]
                    rows.append("\t".join(row_vals))
                sheet_texts.append(f"### Sheet: {ws.title}\n" + "\n".join(rows))
            text = "\n\n".join(sheet_texts)
            meta["source"] = "xlsx"
            return text, meta
        if suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif"):
            with open(path, "rb") as f:
                img = Image.open(io.BytesIO(f.read()))
                img = img.convert("RGB")
            text = ocr_image(img)
            meta["source"] = "image"
            return text, meta
        # Fallback: treat as plain text if possible
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            data = f.read()
        meta["source"] = "plain"
        return data, meta
    except Exception as e:
        logger.exception("Failed to parse file %s: %s", path, e)
        return "", {"name": p.name, "suffix": suffix, "error": str(e)}
