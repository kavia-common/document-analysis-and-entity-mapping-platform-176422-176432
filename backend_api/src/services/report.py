from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Tuple

from openpyxl import Workbook  # type: ignore
from openpyxl.utils import get_column_letter  # type: ignore
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.core.config import get_settings
from src.core.logger import get_logger
from src.db.models import Job, Document, Extraction, Entity, Report

logger = get_logger(__name__)


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(10, length + 2), 60)


# PUBLIC_INTERFACE
async def generate_report_for_job(session: AsyncSession, job_id: int) -> Tuple[str, Dict[str, Any]]:
    """Generate a 5-sheet Excel report for a job and persist report metadata."""
    settings = get_settings()
    out_dir = Path(settings.STORAGE_DIR) / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)

    job = await session.get(Job, job_id)
    if not job:
        raise ValueError("Job not found")

    # Collect data
    docs = (await session.execute(select(Document).where(Document.job_id == job_id))).scalars().all()
    extrs = []
    ents = []
    for d in docs:
        d_extrs = (await session.execute(select(Extraction).where(Extraction.document_id == d.id))).scalars().all()
        extrs.extend(d_extrs)
        for ex in d_extrs:
            ex_ents = (await session.execute(select(Entity).where(Entity.extraction_id == ex.id))).scalars().all()
            ents.extend(ex_ents)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Job ID", job.id])
    ws1.append(["Filename", job.filename])
    ws1.append(["Status", job.status])
    ws1.append(["Documents", len(docs)])
    ws1.append(["Extractions", len(extrs)])
    ws1.append(["Entities", len(ents)])
    _autosize(ws1)

    ws2 = wb.create_sheet("Documents")
    ws2.append(["Doc ID", "Name", "Page Count"])
    for d in docs:
        ws2.append([d.id, d.name, d.page_count or ""])
    _autosize(ws2)

    ws3 = wb.create_sheet("Extractions")
    ws3.append(["Extraction ID", "Document ID", "Model", "Success"])
    for ex in extrs:
        ws3.append([ex.id, ex.document_id, ex.model_name, ex.success])
    _autosize(ws3)

    ws4 = wb.create_sheet("Entities")
    ws4.append(["Entity ID", "Extraction ID", "Type", "Value", "Confidence", "L1", "L2", "L3"])
    for e in ents:
        ws4.append([e.id, e.extraction_id, e.type, e.value, e.confidence, e.l1_id, e.l2_id, e.l3_id])
    _autosize(ws4)

    ws5 = wb.create_sheet("Mappings")
    ws5.append(["Note", "Mappings are applied via mapping API; see DB for details."])
    _autosize(ws5)

    filename = f"job_{job.id}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.xlsx"
    out_path = out_dir / filename
    wb.save(out_path)

    report = Report(job_id=job.id, storage_path=str(out_path), meta={"sheets": 5})
    session.add(report)
    await session.flush()

    logger.info("Generated report for job %s at %s", job.id, out_path)
    return str(out_path), {"sheets": 5}
