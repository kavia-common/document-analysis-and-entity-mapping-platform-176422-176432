from __future__ import annotations

import io
from typing import List

from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

from .parse import ParsedFile


# PUBLIC_INTERFACE
def build_workbook(parsed_files: List[ParsedFile], table_row_cap: int = 1000) -> io.BytesIO:
    """Construct an Excel workbook from parsed files.

    Sheets:
    - Summary: filename, type, size KB, metric, value, status
    - Extracts: filename headers separating up to 10,000 chars of text
    - Tables: Combined tables from CSV/XLSX with file/sheet labels
    - Metadata: content type, parser used, error if any
    """
    wb = Workbook()
    # Default sheet becomes "Summary"
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["File Name", "Type", "Size (KB)", "Metric", "Value", "Status"])

    ws_extracts = wb.create_sheet("Extracts")
    ws_tables = wb.create_sheet("Tables")
    ws_metadata = wb.create_sheet("Metadata")

    # Summary and Metadata
    for pf in parsed_files:
        size_kb = round(pf.size_bytes / 1024.0, 2)
        status = "OK" if not pf.error else f"Error: {pf.error}"
        ws_summary.append([pf.filename, pf.file_type, size_kb, pf.metric_label, pf.metric_value, status])
        ws_metadata.append(["filename", pf.filename])
        ws_metadata.append(["content_type", pf.content_type])
        ws_metadata.append(["parser", pf.parser_used])
        ws_metadata.append(["error", pf.error or ""])
        ws_metadata.append(["---"])

    # Extracts
    for pf in parsed_files:
        if pf.text_excerpt:
            ws_extracts.append([pf.filename])
            ws_extracts.append(["---"])
            # Split large text into lines for readability
            for line in pf.text_excerpt.splitlines():
                ws_extracts.append([line])
            ws_extracts.append([""])
            ws_extracts.append([""])

    # Tables: combine with labels, enforce cap
    ws_tables.append(["Source File", "Sheet Name", "(tabular data below...)"])
    for pf in parsed_files:
        if pf.tables:
            for (sheet_name, df) in pf.tables:
                ws_tables.append([pf.filename, sheet_name])
                # Enforce row cap for each table to keep file small
                df_limited = df.head(table_row_cap)
                # Convert DataFrame to rows and append
                for r in dataframe_to_rows(df_limited, index=False, header=True):
                    ws_tables.append(r)
                ws_tables.append([""])  # spacer

    # Save to in-memory buffer
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
